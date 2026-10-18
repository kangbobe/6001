"""Build the A/B/C English technical draft from completed local artifacts only."""
from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MODEL_ORDER = ['dummy', 'logistic_regression', 'decision_tree', 'random_forest', 'mlp']
DISPLAY = {
    'dummy': 'Dummy prior',
    'logistic_regression': 'Logistic Regression',
    'decision_tree': 'Decision Tree',
    'random_forest': 'Random Forest',
    'mlp': 'MLP',
}
OUTLINE_PATH = ROOT / 'notes/CA6001_ABC_Technical_Draft.md'
DOCX_PATH = ROOT / 'notes/CA6001_ABC_Technical_Draft.docx'


class DraftError(RuntimeError):
    pass


def read_json(path: Path):
    with path.open('r', encoding='utf-8') as stream:
        return json.load(stream)


def require_completed_experiment(experiment: Path):
    completion_path = experiment / 'completion.json'
    if not completion_path.is_file():
        raise DraftError(f'Experiment is still running or incomplete: missing {completion_path}')
    completion = read_json(completion_path)
    if completion.get('status') != 'complete' or completion.get('test_evaluated') is not False:
        raise DraftError('Require status=complete and test_evaluated=false before drafting')
    if completion.get('models') != MODEL_ORDER:
        raise DraftError(f'Unexpected completed model set: {completion.get("models")}')
    protocol = read_json(experiment / 'run_protocol.json')
    if protocol.get('test_evaluated') is not False:
        raise DraftError('Development protocol does not confirm test_evaluated=false')
    artifacts = {}
    for name in MODEL_ORDER:
        model_dir = experiment / name
        metadata = read_json(model_dir / 'model_metadata.json')
        if metadata.get('model') != name or metadata.get('test_evaluated') is not False:
            raise DraftError(f'Missing or unexpected metadata for {name}')
        if metadata.get('refit_partition') != 'train':
            raise DraftError(f'{name} was not refit on train as expected')
        candidates = list(csv.DictReader((model_dir / 'cv_candidates.csv').open(encoding='utf-8', newline='')))
        folds = list(csv.DictReader((model_dir / 'cv_folds.csv').open(encoding='utf-8', newline='')))
        selected_id = str(metadata['config_id'])
        selected = next((row for row in candidates if int(row['config_id']) == int(selected_id)), None)
        if selected is None:
            raise DraftError(f'Selected CV candidate {selected_id} is absent for {name}')
        if not math.isclose(float(selected['mean_average_precision']),
                            float(metadata['cv_average_precision']), rel_tol=1e-10, abs_tol=1e-10):
            raise DraftError(f'CV AP mismatch between candidates and metadata for {name}')
        artifacts[name] = {
            'metadata': metadata,
            'selected_candidate': selected,
            'candidates': candidates,
            'folds': folds,
        }
        if name == 'mlp':
            histories_path = model_dir / 'fit_histories.json'
            if not histories_path.is_file():
                raise DraftError('MLP fit histories are required to describe retries and final warnings')
            histories = read_json(histories_path)
            if len(histories) != len(folds):
                raise DraftError('MLP fit-history count does not match the recorded CV folds')
            histories_by_fold = {(int(row['config_id']), int(row['fold_id'])): row for row in histories}
            if len(histories_by_fold) != len(histories):
                raise DraftError('MLP fit histories contain duplicate config/fold keys')
            for fold in folds:
                history = histories_by_fold.get((int(fold['config_id']), int(fold['fold_id'])))
                if history is None or len(history['attempts']) != int(fold['attempt_count']):
                    raise DraftError('MLP fit-history attempt count does not match the CV fold log')
                if bool(history['attempts'][-1]['convergence_warning']) != _truth(fold['final_convergence_warning']):
                    raise DraftError('MLP fit-history final warning does not match the CV fold log')
                first_warned = bool(history['attempts'][0]['convergence_warning'])
                was_retried = len(history['attempts']) > 1
                if first_warned != was_retried:
                    raise DraftError('MLP retries do not match the fixed warning-triggered retry rule')
            artifacts[name]['histories'] = histories
    return completion, protocol, artifacts


def format_number(value, digits=3):
    return f'{float(value):.{digits}f}'


def _format_param_value(key, value):
    if key == 'weight':
        return 'unweighted' if value is None else str(value)
    if key == 'hidden_layer_sizes':
        layers = tuple(value)
        return str(layers)
    if value is None:
        return 'None'
    return str(value)


def format_params(params: dict) -> str:
    return '; '.join(f'{key}={_format_param_value(key, value)}'
                     for key, value in sorted(params.items()))


def _truth(value):
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {'true', '1', 'yes'}


def _load_project_evidence(experiment):
    completion, protocol, models = require_completed_experiment(experiment)
    eda_dir = ROOT / 'results/eda'
    required_eda = [
        'source_audit.json', 'duplicate_profile_summary.json', 'target_distribution.csv',
        'contact_history_diagnostics.csv', 'fig1_target_distribution.png',
        'fig4_category_subscription_rates.png',
    ]
    missing = [name for name in required_eda if not (eda_dir / name).is_file()]
    if missing:
        raise DraftError(f'Missing train EDA artifacts: {missing}')
    audit = read_json(eda_dir / 'source_audit.json')
    duplicates = read_json(eda_dir / 'duplicate_profile_summary.json')
    target_rows = list(csv.DictReader((eda_dir / 'target_distribution.csv').open(encoding='utf-8', newline='')))
    contact_rows = list(csv.DictReader((eda_dir / 'contact_history_diagnostics.csv').open(encoding='utf-8', newline='')))
    if audit.get('test_used_for_modeling_or_target_association') is not False:
        raise DraftError('Source audit does not certify that test target relationships were not computed')
    if int(audit.get('train', {}).get('rows', -1)) != int(protocol['train_n']):
        raise DraftError('EDA train row count does not match the frozen experiment protocol')
    return completion, protocol, models, audit, duplicates, target_rows, contact_rows, eda_dir


def _cv_table(models):
    rows = []
    for name in MODEL_ORDER:
        artifact = models[name]
        m = artifact['metadata']
        rows.append([
            DISPLAY[name], format_params(m['selected_params']),
            f'{format_number(m["cv_average_precision"])} ± {format_number(m["cv_ap_std"])}',
        ])
    return ['Model', 'Selected train-CV configuration', 'CV AP (mean ± SD)'], rows


def _validation_table(models):
    rows = []
    for name in MODEL_ORDER:
        metrics = models[name]['metadata']['validation_metrics_at_0_5']
        rows.append([
            DISPLAY[name], format_number(metrics['average_precision']),
            format_number(metrics['roc_auc']), format_number(metrics['accuracy']),
            format_number(metrics['f1']), format_number(metrics['precision']),
            format_number(metrics['recall']),
        ])
    return ['Model', 'AP', 'ROC-AUC', 'Accuracy', 'F1', 'Precision', 'Recall'], rows


def _rf_weight_observation(models):
    candidates = models['random_forest']['candidates']
    parsed = [(row, json.loads(row['params'])) for row in candidates]
    matched = []
    for weighted_row, weighted in parsed:
        if weighted.get('weight') != 'balanced':
            continue
        plain = next((row for row, params in parsed
                      if params.get('weight') is None
                      and all(params.get(k) == weighted.get(k) for k in ('max_depth', 'min_samples_leaf'))), None)
        if plain is not None:
            matched.append((weighted_row, plain))
    if matched:
        ap_deltas = [float(w['mean_average_precision'])-float(u['mean_average_precision'])
                     for w, u in matched]
        accuracy_deltas = [float(w['mean_accuracy'])-float(u['mean_accuracy']) for w, u in matched]
        positive = sum(delta > 0 for delta in ap_deltas)
        return (
            f'Across all {len(matched)} matched RF parameter pairs, balanced weighting changed mean '
            f'CV AP by {min(ap_deltas):+.4f} to {max(ap_deltas):+.4f}; '
            f'{positive}/{len(matched)} pairs increased. These small, mixed changes do not establish '
            f'a general ranking benefit from weighting. At 0.5, paired CV accuracy changes ranged '
            f'from {min(accuracy_deltas):+.4f} to {max(accuracy_deltas):+.4f}. Ranking, threshold '
            'performance and probability accuracy are separate objectives; selection used CV AP. '
            'Weighted predict_proba scores are not assumed to be calibrated population probabilities.'
        )
    return (
        'Balanced sample weights change the relative cost of positive-class errors. At the fixed '
        '0.5 threshold, accuracy can move differently from positive-class recall and ranking; '
        'accuracy alone does not establish that a weighted RF is worse. In this experiment, model '
        'selection was based on CV AP, not accuracy.'
    )


def _mlp_observation(models):
    rows = [(row, json.loads(row['params'])) for row in models['mlp']['candidates']]
    by_size = {}
    for row, params in rows:
        size = tuple(params['hidden_layer_sizes'])
        by_size.setdefault(size, []).append(row)
    gaps = {}
    for size, group in by_size.items():
        train_ap = sum(float(r['mean_train_ap']) for r in group) / len(group)
        cv_ap = sum(float(r['mean_average_precision']) for r in group) / len(group)
        gaps[size] = (train_ap, cv_ap, train_ap-cv_ap)
    small, large = (32,), (64, 32)
    selected = models['mlp']['metadata']['config_id']
    selected_row = next(r for r in models['mlp']['candidates'] if int(r['config_id']) == int(selected))
    selected_params = json.loads(selected_row['params'])
    selected_gap = float(selected_row['mean_train_ap']) - float(selected_row['mean_average_precision'])
    cv_fold_rows = [row for row in models['mlp']['folds'] if int(row['config_id']) == int(selected)]
    warning_folds = sum(_truth(row['final_convergence_warning']) for row in cv_fold_rows)
    selected_histories = [row for row in models['mlp']['histories']
                          if int(row['config_id']) == int(selected)]
    if len(selected_histories) != len(cv_fold_rows):
        raise DraftError('Selected MLP fit histories do not match its CV folds')
    selected_retries = sum(len(row['attempts']) > 1 for row in selected_histories)
    retry_count = sum(max(0, len(row['attempts']) - 1) for row in models['mlp']['histories'])
    final_warning_count = sum(bool(row['attempts'][-1]['convergence_warning'])
                              for row in models['mlp']['histories'])
    if any(len(row['attempts']) > 2 for row in models['mlp']['histories']):
        raise DraftError('MLP fit history exceeds the fixed one-retry rule')
    attempts = models['mlp']['metadata']['refit_attempts']
    refit_warning = bool(attempts[-1]['convergence_warning'])
    larger_text = ''
    if small in gaps and large in gaps:
        small_gap, large_gap = gaps[small][2], gaps[large][2]
        if large_gap > small_gap:
            larger_text = (
                f'Averaged over the four tested configurations per architecture, (64, 32) had a train-minus-CV AP gap of '
                f'{format_number(large_gap)} versus {format_number(small_gap)} for (32), '
                'consistent with overfitting in this search. Some larger-network fits still warned; '
                'architecture and optimization effects are not isolated. Absence of a warning means '
                'the configured stopping rule was met, not that a global optimum was found.'
            )
        else:
            larger_text = (
                f'The larger (64, 32) architecture did not have a larger average train-minus-CV '
                f'AP gap than (32) in this search ({format_number(large_gap)} versus '
                f'{format_number(small_gap)}); these records do not support a broader size-based '
                'overfitting claim.'
            )
    return (
        f'Of {len(models["mlp"]["histories"])} MLP CV folds, {retry_count} first attempts required '
        f'the fixed one-retry rule; {final_warning_count} final attempts still warned. '
        f'Selected MLP {format_params(selected_params)} training/CV AP was '
        f'{format_number(selected_row["mean_train_ap"])} / '
        f'{format_number(selected_row["mean_average_precision"])} (gap {format_number(selected_gap)}). '
        f'After retries, its final warning flags were {warning_folds}/{len(cv_fold_rows)} '
        f'({selected_retries} selected fold(s) retried). The train refit finished at '
        f'{attempts[-1]["n_iter"]} iterations with warning={str(refit_warning).lower()} '
        f'({len(attempts)} attempt). {larger_text}'
    )


def _paragraph_blocks(completion, protocol, models, audit, duplicates, target_rows,
                      contact_rows, eda_dir):
    yes_row = next(row for row in target_rows if row['target'] == 'yes')
    train_positive = float(yes_row['share'])
    negative_balance = int(audit['train']['negative_balance'])
    policy_dupes = duplicates['policy_11_profile_pdays_replaced_by_flag']['duplicate_excess_rows']
    prior_unknown = next(int(row['n']) for row in contact_rows
                         if row['diagnostic'] == 'previous_positive_and_outcome_unknown')
    split = (int(protocol['train_n']), int(protocol['validation_n']),
             int(audit['partition_sizes']['test']))
    raw_rows, raw_columns = audit['raw_shape']
    source_input_count = len(duplicates['raw_16_inputs']['columns'])
    version_text = ', '.join(
        f'{name} {version}' for name, version in [
            ('Python', protocol['python']),
            ('scikit-learn', protocol['versions']['scikit-learn']),
            ('NumPy', protocol['versions']['numpy']),
            ('pandas', protocol['versions']['pandas']),
            ('Matplotlib', protocol['versions']['matplotlib']),
            ('joblib', protocol['versions']['joblib']),
        ])
    cv_headers, cv_rows = _cv_table(models)
    val_headers, val_rows = _validation_table(models)
    recommended = DISPLAY.get(completion['cv_recommended_family'], completion['cv_recommended_family'])
    mlp_meta = models['mlp']['metadata']
    mlp_all_warning_folds = sum(bool(row['attempts'][-1]['convergence_warning'])
                                for row in models['mlp']['histories'])
    mlp_total_folds = len(models['mlp']['histories'])
    mlp_retries = sum(len(row['attempts']) > 1 for row in models['mlp']['histories'])

    return [
        ('title', 'CA6001 Applied AI Algorithms Project'),
        ('subtitle', 'A/B/C Technical Draft — Development Evidence Only'),
        ('paragraph', 'Working draft for group integration; student names remain blank. It reports '
                      'development and validation only. D thresholds and final test results remain for D.'),
        ('heading', '1. Introduction'),
        ('paragraph', 'We rank historical records by term-deposit subscription in a retrospective '
                      'pre-campaign screening simulation, comparing four classifiers with a prior '
                      'Dummy baseline. The task estimates neither causal marketing impact nor verified '
                      'historical field availability.'),
        ('heading', '2. Problem Definition'),
        ('paragraph', 'For record i, yᵢ ∈ {0,1} (yes=1); the model produces a subscription score '
                      'sᵢ ∈ [0,1] via predict_proba, with development predictions at sᵢ≥0.5. Scores are '
                      'not guaranteed to be calibrated population probabilities. AP assesses ranking for the minority class '
                      '(scikit-learn developers, n.d.-a). The unit is a record, not a verified customer; '
                      'no customer ID or complete as-of snapshot supports customer-level or future-period claims.'),
        ('heading', '3. Dataset Description'),
        ('paragraph', f'UCI legacy bank-full.csv has {int(raw_rows):,} records, {source_input_count} '
                      f'inputs, and a target ({int(raw_columns)} columns; Moro et al., 2011, 2014). '
                      f'Train/validation/test sizes are {split[0]:,}/{split[1]:,}/{split[2]:,}; train yes rate '
                      f'is {train_positive:.3%}. Eleven retained fields: age, job, marital, education, '
                      'default, balance, housing, loan, pdays, previous, poutcome. Duration, campaign, '
                      'contact, day, and month are excluded.'),
        ('paragraph', f'pdays becomes `has_previous_contact = (pdays >= 0)`; exact days are dropped. '
                      f'Unknowns, negative/extreme balances, and all rows remain. Train has '
                      f'{negative_balance:,} negative balances. Duplicate excess rows (rows minus unique '
                      f'profiles) are {audit["train"]["11_source_feature_duplicates"]:,} for the 11 source fields '
                      f'with exact pdays, and {policy_dupes:,} after replacing pdays with the contact flag. '
                      f'Neither count identifies customers; no rows are removed. There are {prior_unknown} rows with previous>0 but '
                      'unknown poutcome. Figure 1 shows train target distribution.'),
        ('figure', str(eda_dir / 'fig1_target_distribution.png'),
         'Figure 1. Training-partition target distribution (yes = 1).'),
        ('heading', '4. Algorithm Selection and Rationale'),
        ('paragraph', 'Logistic Regression models log-odds as β₀+βᵀx: regularized and economical, '
                      'but linear without added features. A Gini-split Decision Tree captures '
                      'nonlinear rules and interactions and is inspectable, but can be unstable or '
                      'overfit (scikit-learn developers, n.d.-c, n.d.-e).'),
        ('paragraph', 'Random Forest averages bootstrap, feature-sampled trees to reduce variance but '
                      'limits direct interpretability. MLP affine/ReLU layers model nonlinear combinations '
                      'but require scaled numeric inputs and regularization. Dummy prior is the baseline; '
                      'boosting and time/customer-grouped evaluation remain unrun alternatives '
                      '(scikit-learn developers, n.d.-d, n.d.-f, n.d.-g).'),
        ('heading', '5. Methodology'),
        ('paragraph', f'The fixed stratified 60/20/20 split uses seed 42. Only train ({split[0]:,}) '
                      'enters shuffled 3-fold StratifiedKFold CV. Prespecified searches maximize mean AP; '
                      'lower config_id breaks ties. AP is non-interpolated average_precision_score, not '
                      'trapezoidal PR-AUC (scikit-learn developers, n.d.-a). Winners are refit on train '
                      'and scored on validation for diagnosis and D’s threshold choice, not further tuning. '
                      'No test predictions or test-performance analysis are reported.'),
        ('paragraph', 'A shared pipeline derives the prior-contact flag and one-hot encodes categories. '
                      'The literal value "unknown" is retained as a category. Categories not seen during '
                      'the current fit are encoded as an all-zero block at transform time; records remain. '
                      'Scale numeric inputs only for Logistic Regression '
                      'and MLP. Balanced weights are computed per fit fold; scoring is unweighted. No '
                      'SMOTE, global scaling, clipping, or deduplication; development threshold is 0.5 '
                      '(scikit-learn developers, n.d.-b).'),
        ('paragraph', 'Searches contain 8 LR, 12 DT, 8 RF and 8 MLP configurations (36 total; '
                      '108 base CV fits). LR C={0.01,0.1,1,10}; DT depth={4,8,None} and leaf={100,20}; '
                      'RF depth={10,None} and leaf={20,5}; MLP layers={(32,),(64,32)} and '
                      'alpha={0.01,0.0001}. Each crosses unweighted/balanced fits. RF uses 200 trees, '
                      'sqrt feature subsampling and bootstrap. MLP uses Adam/ReLU, learning rate=0.001, '
                      'batch=256, early_stopping=False, tol=0.0001 and n_iter_no_change=20. '
                      'Every MLP fit that reaches 300 iterations with a ConvergenceWarning restarts '
                      'once from scratch with the same seed and a 600-iteration cap; final warnings '
                      'remain in the record. Mean ± SD uses three folds and sample SD (ddof=1), '
                      'not a confidence interval. Full configurations are in search_configs.json.'),
        ('paragraph', f'Run versions: {version_text} (Python Software Foundation, n.d.; NumPy developers, n.d.; '
                      'pandas development team, n.d.; Matplotlib developers, n.d.; Joblib developers, n.d.). '
                      'Codex/ChatGPT Pro assisted drafting and software '
                      'development (OpenAI, n.d.); reported metrics come from local run artifacts.'),
        ('heading', '6. Implementation'),
        ('paragraph', 'The data module validates source and splits; preprocessing stays in the pipeline. '
                      'Training logs fold candidates, refits winners on train, reloads pipelines, and '
                      'checks validation probabilities. Core functions: read_raw, load_partition, '
                      'audit_data, build_preprocessor, score_fold, train_family, run_experiment.'),
        ('code', 'for fit_ids, score_ids in folds:\n'
                 '    fitted, history = fit_candidate(name, config, X.iloc[fit_ids], y.iloc[fit_ids])\n'
                 '    p_yes = fitted.predict_proba(X.iloc[score_ids])[:, 1]\n'
                 '    fold_ap = average_precision_score(y.iloc[score_ids], p_yes)'),
        ('paragraph', f'Across {mlp_total_folds} MLP CV folds, {mlp_retries} required one retry and '
                      f'{mlp_all_warning_folds} final attempts warned. See Section 8 for selected-fold details.'),
        ('heading', '7. Results and Evaluation — Development Only'),
        ('paragraph', f'The completed run recommends {recommended} by CV AP for development. Table 1 '
                      'shows selected configurations and CV AP mean ± SD; Table 2 gives validation '
                      'metrics at 0.5. These are development results, not a final test conclusion.'),
        ('table', *_cv_table(models)),
        ('table', *_validation_table(models)),
        ('paragraph', _rf_weight_observation(models)),
        ('figure', str(eda_dir / 'fig4_category_subscription_rates.png'),
         'Figure 2. Training-only subscription rates by prior campaign outcome and job; labels show group n.'),
        ('heading', '8. Discussion & Insights'),
        ('paragraph', _mlp_observation(models)),
        ('paragraph', 'Historical row-split results describe record-level generalization only: no '
                      'customer ID supports customer-disjoint evaluation, and the prediction-time '
                      'snapshot is incomplete. They do not establish future-campaign performance or '
                      'causal effects. AP measures ranking; Table 2 precision, recall, F1, and accuracy '
                      'use 0.5. Figure 2 is descriptive, and data vintage or population bias may shape '
                      'subgroup patterns.'),
        ('heading', '9. Conclusion — D to Complete'),
        ('paragraph', 'D will complete the final conclusion after freezing thresholds and running the '
                      'authorized one-shot test evaluation, integrating final metrics, error and '
                      'importance analyses. No final winner or test result is asserted here.'),
        ('heading', '10. References'),
        ('reference', 'Joblib developers. (n.d.). *joblib* (Version 1.6.0) [Computer software]. https://joblib.readthedocs.io/'),
        ('reference', 'Matplotlib developers. (n.d.). *Matplotlib* (Version 3.11.2) [Computer software]. https://matplotlib.org/'),
        ('reference', 'Moro, S., Laureano, R., & Cortez, P. (2011). Using data mining for bank direct marketing: An application of the CRISP-DM methodology. In P. Novais et al. (Eds.), *Proceedings of the European Simulation and Modelling Conference (ESM’2011)* (pp. 117–121). EUROSIS. https://hdl.handle.net/1822/14838'),
        ('reference', 'Moro, S., Rita, P., & Cortez, P. (2014). *Bank Marketing* [Data set]. UCI Machine Learning Repository. https://doi.org/10.24432/C5K306'),
        ('reference', 'NumPy developers. (n.d.). *NumPy* (Version 2.4.6) [Computer software]. https://numpy.org/'),
        ('reference', 'OpenAI. (n.d.). *ChatGPT Pro and Codex* [AI-assisted drafting and software-development tools]. https://openai.com/'),
        ('reference', 'pandas development team. (n.d.). *pandas* (Version 3.0.3) [Computer software]. https://pandas.pydata.org/'),
        ('reference', 'Python Software Foundation. (n.d.). *Python* (Version 3.12.10) [Computer software]. https://www.python.org/'),
        ('reference', 'scikit-learn developers. (n.d.-a). *Average precision score*. scikit-learn 1.8.0 API documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.metrics.average_precision_score.html'),
        ('reference', 'scikit-learn developers. (n.d.-b). *Common pitfalls and recommended practices*. scikit-learn 1.8.0 user guide. https://scikit-learn.org/1.8/common_pitfalls.html'),
        ('reference', 'scikit-learn developers. (n.d.-c). *DecisionTreeClassifier*. scikit-learn 1.8.0 API documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.tree.DecisionTreeClassifier.html'),
        ('reference', 'scikit-learn developers. (n.d.-d). *DummyClassifier*. scikit-learn 1.8.0 API documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.dummy.DummyClassifier.html'),
        ('reference', 'scikit-learn developers. (n.d.-e). *LogisticRegression*. scikit-learn 1.8.0 API documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.linear_model.LogisticRegression.html'),
        ('reference', 'scikit-learn developers. (n.d.-f). *MLPClassifier*. scikit-learn 1.8.0 API documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.neural_network.MLPClassifier.html'),
        ('reference', 'scikit-learn developers. (n.d.-g). *RandomForestClassifier*. scikit-learn 1.8.0 API documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.ensemble.RandomForestClassifier.html'),
    ]


def render_markdown(blocks) -> str:
    lines = []
    table_number = 0
    for block in blocks:
        kind = block[0]
        if kind == 'title':
            lines += [f'# {block[1]}', '']
        elif kind == 'subtitle':
            lines += [f'## {block[1]}', '']
        elif kind == 'heading':
            lines += [f'## {block[1]}', '']
        elif kind == 'paragraph':
            lines += [block[1], '']
        elif kind == 'code':
            lines += ['```python', block[1], '```', '']
        elif kind == 'figure':
            path = Path(block[1])
            rel_path = Path('..') / 'results' / 'eda' / path.name
            lines += [f'![{block[2]}]({rel_path.as_posix()})', '', f'*{block[2]}*', '']
        elif kind == 'table':
            table_number += 1
            headers, rows = block[1], block[2]
            lines += [f'**Table {table_number}. {"Selected train-CV configurations" if table_number == 1 else "Validation metrics at threshold 0.5"}.**', '']
            lines += ['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |']
            lines += ['| ' + ' | '.join(str(cell).replace('|', '\\|') for cell in row) + ' |' for row in rows]
            lines += ['']
        elif kind == 'reference':
            lines += [block[1], '']
    return '\n'.join(lines).rstrip() + '\n'


def _add_reference_text(paragraph, text):
    parts = re.split(r'(\*[^*]+\*)', text)
    for part in parts:
        if part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part:
            paragraph.add_run(part)


def render_docx(blocks, path: Path):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.62)
    section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.72)
    section.right_margin = Inches(.72)
    normal = document.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    for heading_style in ('Heading 1', 'Heading 2'):
        document.styles[heading_style].font.name = 'Aptos Display'
    expected_tables = []
    expected_paragraphs = []
    image_count = 0
    table_number = 0
    for block in blocks:
        kind = block[0]
        if kind == 'title':
            p = document.add_heading(block[1], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            expected_paragraphs.append(block[1])
        elif kind == 'subtitle':
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(block[1])
            run.bold = True
            run.font.size = Pt(13)
            expected_paragraphs.append(block[1])
        elif kind == 'heading':
            document.add_heading(block[1], level=1)
            expected_paragraphs.append(block[1])
        elif kind == 'paragraph':
            p = document.add_paragraph(block[1])
            expected_paragraphs.append(block[1])
        elif kind == 'code':
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(.2)
            p.paragraph_format.right_indent = Inches(.2)
            run = p.add_run(block[1])
            run.font.name = 'Consolas'
            run.font.size = Pt(8)
            expected_paragraphs.append(block[1])
        elif kind == 'figure':
            figure_path = Path(block[1])
            if not figure_path.is_file():
                raise DraftError(f'Missing report figure: {figure_path}')
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(figure_path), width=Inches(6.2))
            cap = document.add_paragraph(block[2])
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            image_count += 1
            expected_paragraphs.append(block[2])
        elif kind == 'table':
            table_number += 1
            headers, rows = block[1], block[2]
            table = document.add_table(rows=1, cols=len(headers))
            table.style = 'Light Shading Accent 1'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, header in enumerate(headers):
                table.rows[0].cells[i].text = header
            for row in rows:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value)
            for row in table.rows:
                for cell in row.cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.name = 'Aptos'
                            run.font.size = Pt(7.5)
            expected_tables.append([[str(value) for value in headers]] +
                                   [[str(value) for value in row] for row in rows])
            caption = 'Selected train-CV configurations' if table_number == 1 else 'Validation metrics at threshold 0.5'
            p = document.add_paragraph(f'Table {table_number}. {caption}.')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].italic = True
            expected_paragraphs.append(p.text)
        elif kind == 'reference':
            p = document.add_paragraph()
            _add_reference_text(p, block[1])
            p.paragraph_format.left_indent = Inches(.25)
            p.paragraph_format.first_line_indent = Inches(-.25)
            p.paragraph_format.space_after = Pt(3)
            expected_paragraphs.append(re.sub(r'\*([^*]+)\*', r'\1', block[1]))
    document.save(path)
    readback = Document(path)
    actual_tables = [[[cell.text for cell in row.cells] for row in table.rows]
                     for table in readback.tables]
    if actual_tables != expected_tables:
        raise DraftError('DOCX table readback mismatch; one or more table cells changed')
    if len(readback.inline_shapes) != image_count or image_count != 2:
        raise DraftError(f'DOCX should contain exactly 2 selected figures; found {len(readback.inline_shapes)}')
    actual_paragraphs = [p.text for p in readback.paragraphs]
    missing = [text for text in expected_paragraphs if text not in actual_paragraphs]
    if missing:
        raise DraftError(f'DOCX paragraph readback mismatch: {missing[:2]}')
    return {'paragraphs': len(actual_paragraphs), 'tables': len(actual_tables), 'figures': len(readback.inline_shapes)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--experiment', type=Path, default=ROOT / 'results/experiment')
    parser.add_argument('--output-md', type=Path, default=OUTLINE_PATH)
    parser.add_argument('--output-docx', type=Path, default=DOCX_PATH)
    args = parser.parse_args(argv)

    # The remaining evidence is fixed at the repository's train-only EDA outputs.
    evidence = _load_project_evidence(args.experiment)
    completion, protocol, models, audit, duplicates, target_rows, contact_rows, eda_dir = evidence
    blocks = _paragraph_blocks(completion, protocol, models, audit, duplicates,
                               target_rows, contact_rows, eda_dir)
    markdown = render_markdown(blocks)
    args.output_md.parent.mkdir(parents=True, exist_ok=True)
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    args.output_md.write_text(markdown, encoding='utf-8')
    docx_check = render_docx(blocks, args.output_docx)
    words = len(re.findall(r"\b[\w’'-]+\b", re.sub(r'!\[[^]]*\]\([^)]*\)', '', markdown)))
    print(f'Wrote {args.output_md} and {args.output_docx}')
    print(f'Draft word count (including tables and references): {words}')
    print(f'DOCX readback: {docx_check}')
    print('Final test was not evaluated; Section 9 remains assigned to D.')
    return args.output_md, args.output_docx


if __name__ == '__main__':
    main()
